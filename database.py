import os
import json
import sqlite3
import hashlib
from threading import Lock

import openpyxl
from sqlalchemy import (
    JSON, Column, DateTime, Float, Integer, MetaData, String, Table,
    UniqueConstraint, create_engine, delete, func, insert, select,
)
from sqlalchemy.dialects.postgresql import insert as postgres_insert
from sqlalchemy.dialects.sqlite import insert as sqlite_insert
from sqlalchemy.engine import URL, Engine

from config import DATABASE_FILE, UPLOADS_FILE, EXCEL_FILE
from utils import safe_float, _row_dict, _find_field, _source_color, _clean_text

DATABASE_LOCK = Lock()
ENGINE_LOCK = Lock()

_CACHED_ENGINE = None
_CACHED_ENGINE_KEY = None

DATABASE_METADATA = MetaData()
PROJECTS_TABLE = Table(
    "projects",
    DATABASE_METADATA,
    Column("id", Integer, primary_key=True, autoincrement=True),
    Column("source_type", String, nullable=False),
    Column("source_name", String, nullable=False),
    Column("source_key", String, nullable=False),
    Column("project_name", String, nullable=False),
    Column("external_id", String),
    Column("latitude", Float, nullable=False),
    Column("longitude", Float, nullable=False),
    Column("category", String, nullable=False, default=""),
    Column("fields_json", JSON, nullable=False),
    Column("marker_color", String, nullable=False),
    Column("created_at", DateTime(timezone=True), server_default=func.now()),
    UniqueConstraint("source_type", "source_name", "source_key"),
)
APP_METADATA_TABLE = Table(
    "app_metadata",
    DATABASE_METADATA,
    Column("key", String, primary_key=True),
    Column("value", String, nullable=False),
)

def _database_engine() -> Engine:
    global _CACHED_ENGINE, _CACHED_ENGINE_KEY

    database_url = os.environ.get("DATABASE_URL", "").strip()

    if database_url:
        if database_url.startswith("postgres://"):
            database_url = "postgresql+psycopg://" + database_url[len("postgres://"):]
        elif database_url.startswith("postgresql://"):
            database_url = "postgresql+psycopg://" + database_url[len("postgresql://"):]

        if "sslmode=" not in database_url.lower():
            database_url += "&sslmode=require" if "?" in database_url else "?sslmode=require"

        engine_key = database_url
        engine_url = database_url
        connect_args = {"connect_timeout": 10}
    else:
        local_path = os.path.abspath(DATABASE_FILE)
        engine_key = f"sqlite:{local_path}"
        engine_url = URL.create("sqlite", database=local_path)
        connect_args = {"check_same_thread": False, "timeout": 30}

    with ENGINE_LOCK:
        if _CACHED_ENGINE is not None and _CACHED_ENGINE_KEY == engine_key:
            return _CACHED_ENGINE

        if _CACHED_ENGINE is not None:
            _CACHED_ENGINE.dispose()

        _CACHED_ENGINE = create_engine(
            engine_url,
            connect_args=connect_args,
            pool_pre_ping=bool(database_url),
        )
        _CACHED_ENGINE_KEY = engine_key
        return _CACHED_ENGINE

def database_backend():
    return "Tiger Cloud PostgreSQL" if os.environ.get("DATABASE_URL", "").strip() else "Local SQLite"

def _database_connection():
    return _database_engine().connect()

def _project_key(project):
    project_id = _clean_text(project.get("id")).casefold()
    if project_id:
        return "id", project_id
    return (
        "name-location",
        _clean_text(project.get("name")).casefold(),
        round(float(project["lat"]), 5),
        round(float(project["lon"]), 5),
    )

def _insert_database_project(connection, project):
    fields = project.get("fields") or {}
    source_type = project.get("source_type", "upload")
    source_name = project.get("source_name") or project.get("sheet") or "Imported"
    source_key = project.get("source_key")

    if not source_key:
        key_data = json.dumps(_project_key(project), ensure_ascii=False, default=str)
        source_key = hashlib.sha256(key_data.encode("utf-8")).hexdigest()

    values = {
        "source_type": source_type,
        "source_name": source_name,
        "source_key": source_key,
        "project_name": _clean_text(project.get("name"), 300) or "Unnamed project",
        "external_id": _clean_text(project.get("id"), 120) or None,
        "latitude": float(project["lat"]),
        "longitude": float(project["lon"]),
        "category": _clean_text(project.get("category"), 120),
        "fields_json": json.loads(json.dumps(fields, ensure_ascii=False, default=str)),
        "marker_color": project.get("color") or _source_color(source_name),
    }
    insert_function = (
        postgres_insert
        if connection.dialect.name == "postgresql"
        else sqlite_insert
    )
    statement = insert_function(PROJECTS_TABLE).values(**values).on_conflict_do_nothing(
        index_elements=[
            PROJECTS_TABLE.c.source_type,
            PROJECTS_TABLE.c.source_name,
            PROJECTS_TABLE.c.source_key,
        ]
    )
    return connection.execute(statement).rowcount == 1

def _project_from_database_row(row):
    fields = row["fields_json"]
    if isinstance(fields, str):
        fields = json.loads(fields)
    return {
        # The database's own primary key -- the only value that's always
        # unique and stable, even for uploads with no external project_id.
        # This is what delete_project() below expects, and what the
        # frontend needs to send back when the user deletes a project.
        "db_id": row["id"],
        "source_type": row["source_type"],
        "source_name": row["source_name"],
        "sheet": row["source_name"],
        "id": row["external_id"],
        "name": row["project_name"],
        "category": row["category"],
        "lat": row["latitude"],
        "lon": row["longitude"],
        "fields": fields,
        "color": row["marker_color"],
    }

def initialize_database():
    with DATABASE_LOCK:
        engine = _database_engine()
        DATABASE_METADATA.create_all(engine)

        with engine.begin() as connection:
            if engine.dialect.name == "postgresql":
                local_migration = connection.execute(
                    select(APP_METADATA_TABLE.c.value).where(
                        APP_METADATA_TABLE.c.key == "local_uploads_migrated"
                    )
                ).scalar_one_or_none()
                if not local_migration:
                    if os.path.isfile(DATABASE_FILE):
                        try:
                            local_connection = sqlite3.connect(DATABASE_FILE)
                            local_connection.row_factory = sqlite3.Row
                            local_projects = local_connection.execute(
                                "SELECT * FROM projects WHERE source_type = 'upload'"
                            ).fetchall()
                        except sqlite3.Error:
                            local_projects = []
                        finally:
                            if "local_connection" in locals():
                                local_connection.close()

                        for row in local_projects:
                            _insert_database_project(connection, _project_from_database_row(row))

                    connection.execute(
                        insert(APP_METADATA_TABLE).values(
                            key="local_uploads_migrated", value="1"
                        )
                    )

            legacy_migrated = connection.execute(
                select(APP_METADATA_TABLE.c.value).where(
                    APP_METADATA_TABLE.c.key == "legacy_uploads_migrated"
                )
            ).scalar_one_or_none()
            if not legacy_migrated:
                if os.path.isfile(UPLOADS_FILE):
                    with open(UPLOADS_FILE, "r", encoding="utf-8") as uploads_file:
                        legacy_projects = json.load(uploads_file)

                    if not isinstance(legacy_projects, list):
                        raise ValueError("The saved upload data is not a project list.")

                    for project in legacy_projects:
                        if not isinstance(project, dict):
                            continue
                        migrated = dict(project)
                        migrated["source_type"] = "upload"
                        migrated["source_name"] = migrated.get("source_name") or "Company uploads"
                        migrated["fields"] = migrated.get("fields") or {
                            "Imported details": " ".join(
                                str(line) for line in migrated.get("popup_lines", [])
                            )
                        }
                        _insert_database_project(connection, migrated)

                connection.execute(
                    insert(APP_METADATA_TABLE).values(
                        key="legacy_uploads_migrated", value="1"
                    )
                )

def load_workbook_projects(path):
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    projects = []
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            header = next(rows, ())
            if not header:
                continue

            for row_number, row in enumerate(rows, start=2):
                record = _row_dict(header, row)
                latitude = safe_float(_find_field(record, "latitude"))
                longitude = safe_float(_find_field(record, "longitude"))

                if (
                    latitude is None
                    or longitude is None
                    or not -90 <= latitude <= 90
                    or not -180 <= longitude <= 180
                ):
                    continue

                name = _clean_text(_find_field(record, "name"), 300)
                identifier = _find_field(record, "identifier")
                category = _find_field(record, "category")
                projects.append({
                    "source_type": "workbook",
                    "source_name": worksheet.title,
                    "source_key": f"{worksheet.title}:{row_number}",
                    "sheet": worksheet.title,
                    "id": _clean_text(identifier, 120) or None,
                    "name": name or f"Unnamed project (row {row_number})",
                    "category": _clean_text(category, 120),
                    "lat": latitude,
                    "lon": longitude,
                    "fields": record,
                    "color": _source_color(worksheet.title),
                })
    finally:
        workbook.close()
    return projects

def _metadata_value(connection, key):
    return connection.execute(
        select(APP_METADATA_TABLE.c.value).where(APP_METADATA_TABLE.c.key == key)
    ).scalar_one_or_none()

def _set_metadata(connection, key, value):
    insert_function = (
        postgres_insert
        if connection.dialect.name == "postgresql"
        else sqlite_insert
    )
    statement = insert_function(APP_METADATA_TABLE).values(key=key, value=value)
    statement = statement.on_conflict_do_update(
        index_elements=[APP_METADATA_TABLE.c.key],
        set_={"value": value},
    )
    connection.execute(statement)

def sync_workbook_database(path):
    digest = hashlib.sha256()
    with open(path, "rb") as workbook_file:
        for chunk in iter(lambda: workbook_file.read(1024 * 1024), b""):
            digest.update(chunk)
    workbook_hash = digest.hexdigest()

    initialize_database()
    with DATABASE_LOCK:
        with _database_engine().connect() as connection:
            stored_hash = _metadata_value(connection, "workbook_sha256")
            if stored_hash == workbook_hash:
                return

    projects = load_workbook_projects(path)

    with DATABASE_LOCK:
        with _database_engine().begin() as connection:
            stored_hash = _metadata_value(connection, "workbook_sha256")
            if stored_hash == workbook_hash:
                return

            connection.execute(
                delete(PROJECTS_TABLE).where(PROJECTS_TABLE.c.source_type == "workbook")
            )
            for project in projects:
                _insert_database_project(connection, project)
            _set_metadata(connection, "workbook_sha256", workbook_hash)

def _load_database_projects(source_type=None):
    statement = select(PROJECTS_TABLE)
    if source_type is not None:
        statement = statement.where(PROJECTS_TABLE.c.source_type == source_type)
    statement = statement.order_by(
        PROJECTS_TABLE.c.source_type,
        PROJECTS_TABLE.c.source_name,
        PROJECTS_TABLE.c.project_name,
    )

    with DATABASE_LOCK:
        with _database_engine().connect() as connection:
            rows = connection.execute(statement).mappings().all()
            return [_project_from_database_row(row) for row in rows]

def load_uploaded_projects():
    return _load_database_projects("upload")

def load_all_projects():
    sync_workbook_database(EXCEL_FILE)
    return _load_database_projects()

def save_uploaded_projects(projects):
    initialize_database()
    with DATABASE_LOCK:
        with _database_engine().begin() as connection:
            added = sum(_insert_database_project(connection, project) for project in projects)
    return added, len(load_uploaded_projects())

def clear_uploaded_projects():
    """Deletes every uploaded project. Kept for a "clear all uploads"
    action -- use delete_project() for removing a single project."""
    initialize_database()
    with DATABASE_LOCK:
        with _database_engine().begin() as connection:
            connection.execute(
                delete(PROJECTS_TABLE).where(PROJECTS_TABLE.c.source_type == "upload")
            )


def delete_project(db_id):
    """Deletes exactly one project by its database primary key (the
    "db_id" field from _project_from_database_row / the API response) --
    not by name, external_id, or any other value that might collide or be
    missing. Returns True if a row was actually deleted, False if db_id
    didn't match anything (already gone, or never existed).

    This is what actually frees up the project's source_key so a later
    reimport of the same document isn't silently skipped by
    on_conflict_do_nothing() in _insert_database_project() -- if the row
    still exists, a reimport with the same id/name/coordinates hashes to
    the same source_key and gets treated as a duplicate.
    """
    initialize_database()
    with DATABASE_LOCK:
        with _database_engine().begin() as connection:
            result = connection.execute(
                delete(PROJECTS_TABLE).where(PROJECTS_TABLE.c.id == db_id)
            )
            return result.rowcount > 0